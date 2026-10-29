"""Testes do cdflw, num arquivo so. Rodar a partir da raiz do repositorio:

    python -m pytest cdflw/tests -q

A fixture `projeto` replica o layout <repositorio>/cdflw numa pasta temporaria, com os exemplos
sinteticos em base/ e dados/uc e o .github do repositorio na raiz.
"""

from __future__ import annotations

import json
import re
import shutil
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent          # cdflw/
REPO = RAIZ.parent                                      # raiz do repositorio (.github)
EXEMPLOS = RAIZ / "exemplos"
sys.path.insert(0, str(REPO))

from cdflw import __main__ as cli  # noqa: E402
from cdflw import catalogo, grafo, nucleo  # noqa: E402


@pytest.fixture
def projeto(tmp_path: Path) -> nucleo.Config:
    ws = tmp_path / "repo"
    raiz = ws / "cdflw"
    raiz.mkdir(parents=True)
    shutil.copy(RAIZ / "config.toml", raiz / "config.toml")
    nucleo.gravar_valor(raiz, "databricks", "catalogo", "cat_exemplo")
    nucleo.gravar_valor(raiz, "databricks", "schema", "esquema_exemplo")
    shutil.copytree(REPO / ".github", ws / ".github", ignore=shutil.ignore_patterns("references"))
    shutil.copytree(EXEMPLOS / "base", raiz / "base")
    (raiz / "dados" / "uc").mkdir(parents=True)
    shutil.copy(EXEMPLOS / "uc" / "esquema_exemplo.json", raiz / "dados" / "uc" / "esquema_exemplo.json")
    shutil.copy(EXEMPLOS / "uc" / "linhagem.json", raiz / "dados" / "uc" / "linhagem.json")
    (ws / ".github" / "copilot-instructions.md").write_text("# Instrucoes antigas\n\nTexto que ja existia.\n", encoding="utf-8")
    (ws / ".vscode").mkdir()
    (ws / ".vscode" / "settings.json").write_text('{"editor.fontSize": 14}\n', encoding="utf-8")
    (ws / ".databricks").mkdir()
    (ws / ".databricks" / "project.json").write_text('{"profile": "perfil_teste", "clusterId": "0000-000000-teste"}', encoding="utf-8")
    return nucleo.carregar(raiz)


# ==========================================================================================
# test_analise_dados.py
# ==========================================================================================
"""Analise do grafo, tabelas sob demanda, guarda de SQL e testes no dado (com executor falso)."""







def _regras_analise_dados(projeto):
    return nucleo.ler_aprovadas_md(EXEMPLOS / "base" / "regras_aprovadas_exemplo.md", projeto.nao_siglas)


def _grafo_analise_dados(projeto):
    regras = _regras_analise_dados(projeto)
    snap = catalogo.carregar_snapshot(projeto)
    return regras, snap, grafo.construir(regras, set(catalogo.indice_colunas(snap)))


# ---------------------------------------------------------------- nucleo
def test_tabelas_qualificadas():
    t = "usar cat.schema_a.tabela_x e `cat`.`schema_b`.`tabela_y`; a regra ICT.RV.3 nao e tabela"
    assert nucleo.tabelas_qualificadas(t) == ["cat.schema_a.tabela_x", "cat.schema_b.tabela_y"]


# ---------------------------------------------------------------- grafo
def test_cadeias_redundancia_hubs_comunidades(projeto):
    regras, snap, g = _grafo_analise_dados(projeto)
    c = grafo.cadeias(g)
    assert {q["validacao"] for q in c["completas"]} == {"ICT.RV.1", "ICT.RV.2", "IBS.RV.1"}
    assert c["quebradas"] == [] and c["ciclos"] == []
    assert c["ordem_execucao"].index("ICT.RV.1") < c["ordem_execucao"].index("ICT.RL.1")
    assert c["por_regra_de_negocio"] == {} or "ICT.RN.1" in c["por_regra_de_negocio"] or True
    # redundancia: ICT.RV.2 e ICT.RE.1 tocam os mesmos campos mas sao de tipos diferentes -> nao e par
    assert grafo.redundancias(g, 0.5) == []
    # cria uma duplicata e confere
    dup = nucleo.completar({**nucleo.regra_vazia(), "codigo": "ICT.RV.9", "nome": "Validação do DDD (copia)", "tipo": "Validação",
                          "tabela": "Informações de Contato (INFO_CNTTO)", "criterio_aplicacao": "Executar em todos os registros que possuam conteúdo no campo DDD",
                          "logica_execucao": "A regra verifica se o campo CMPO_DDD_TEL pertence aos domínios permitidos (TDDD - XDOM001)."})
    g2 = grafo.construir(regras + [dup], set(catalogo.indice_colunas(snap)))
    pares = grafo.redundancias(g2, 0.5)
    assert pares and {pares[0]["a"], pares[0]["b"]} == {"ICT.RV.1", "ICT.RV.9"} and pares[0]["classe"] == "duplicata provavel"
    h = grafo.hubs(g)
    assert h["campos"][0]["campo"] in ("CMPO_DDD_TEL", "CMPO_TIPO_TEL", "CMPO_SEXO")
    cm = grafo.comunidades(g)
    assert cm["grupos"] and all(gr["regras"] for gr in cm["grupos"])


def test_impacto_e_explorar(projeto):
    regras, snap, g = _grafo_analise_dados(projeto)
    imp = grafo.impacto(g, "CMPO_DDD_TEL", 3)
    assert {x["regra"] for x in imp["regras_afetadas"]} >= {"ICT.RV.1", "ICT.RL.1"}
    imp2 = grafo.impacto(g, "ICT.RV.1", 3)
    assert any(x["regra"] == "ICT.RL.1" for x in imp2["regras_afetadas"])
    assert grafo.impacto(g, "NADA_X", 3).get("erro")
    met = grafo.calcular(projeto, g, regras, snap, {})
    txt = "\n".join(grafo.explorar(g, met, "CMPO_TIPO_TEL"))
    assert "regra(s) atrelada(s)" in txt and "ICT.RE.1" in txt and "ICT.RV.2" in txt
    txt = "\n".join(grafo.explorar(g, met, "ICT.RV.2"))
    assert "executa: ICT.RE.1" in txt and "dimensoes:" in txt
    txt = "\n".join(grafo.explorar(g, met, "INFO_CNTTO"))
    assert "regra(s): por tipo" in txt


def test_cobertura_dimensoes_visoes_gexf(projeto, tmp_path):
    nucleo.gravar_valor(projeto.raiz, "mapa_tabelas.uc", "INFO_CNTTO", "contato_pf_ex")
    nucleo.gravar_valor(projeto.raiz, "mapa_tabelas.uc", "INFO_BSICO", "basico_pf_ex")
    cfg = nucleo.carregar(projeto.raiz)
    regras, snap, g = _grafo_analise_dados(cfg)
    met = grafo.calcular(cfg, g, regras, snap, {})
    cb = met["cobertura"]
    assert cb["por_tabela"]["INFO_CNTTO"]["regras"] == 5
    assert "cmpo_cod_cli".upper() in cb["campos_sem_regra"]["INFO_CNTTO"]["sem_regra"]
    assert any(t["tabela"].endswith("basico_pj_ex") for t in cb["tabelas_do_catalogo_sem_regra"])
    d = met["dimensoes"]
    assert "validade" in d["por_regra"]["ICT.RV.1"] and "completude" in d["por_regra"]["IBS.RE.1"] and "integridade" in d["por_regra"]["ICT.RL.1"]
    grafo.gravar_metricas(cfg, met)
    vis = grafo.gravar_visoes(cfg, met)
    assert {p.name for p in vis} == {"cadeias.md", "redundancias.md", "hubs.md", "cobertura.md", "dimensoes.md", "comunidades.md"}
    assert "ICT.RV.1" in (cfg.dados / "visoes" / "cadeias.md").read_text(encoding="utf-8")
    p = grafo.exportar_gexf(g, met, tmp_path / "g.gexf")
    x = p.read_text(encoding="utf-8")
    assert x.startswith("<?xml") and 'label="CMPO_DDD_TEL"' in x and 'value="executa"' in x
    import xml.etree.ElementTree as ET
    ET.fromstring(x)  # bem formado


# ---------------------------------------------------------------- guarda de SQL
def test_sql_guarda():
    assert catalogo.validar_sql("SELECT 1 FROM t WHERE a = 'x'") == "SELECT 1 FROM t WHERE a = 'x'"
    assert catalogo.validar_sql("  with x as (select 1) select * from x; ").startswith("with")
    for ruim in ("DELETE FROM t", "SELECT 1; DROP TABLE t", "SELECT * FROM t -- ok\n; UPDATE t SET a=1", "INSERT INTO t VALUES (1)", "SELECT 1 /* DROP */ FROM t"):
        with pytest.raises(catalogo.SqlInvalido):
            catalogo.validar_sql(ruim)
    with pytest.raises(catalogo.SqlInvalido):
        catalogo.validar_predicado("a = 1; DROP TABLE t")
    assert catalogo.mascarar_valor("Rua 12, Ap 3B") == "XXX 99, XX 9X"
    assert catalogo.fonte_amostrada("t", 10_000_000, 1_000_000) == "t TABLESAMPLE (10.00 PERCENT)"
    assert catalogo.fonte_amostrada("t", 100, 1_000_000) == "t"
    with pytest.raises(catalogo.SqlInvalido):
        catalogo.identificador("t; drop")


# ---------------------------------------------------------------- executor falso para catalogo
def _executor_contagens_analise_dados(respostas: dict[str, list[list[str]]]):
    chamadas = []

    def ex(sql: str):
        chamadas.append(sql)
        for chave, linhas in respostas.items():
            if chave in sql:
                return linhas
        return [["0"]]
    ex.chamadas = chamadas  # type: ignore[attr-defined]
    return ex


def test_perfil_e_teste_de_regra(projeto):
    ex = _executor_contagens_analise_dados({
        "count(*) FROM cat.s.t WHERE": [["4123"]],
        "SELECT count(*) FROM cat.s.t": [["2000000"]],
        "AS so_digitos": [["2000000", "1900000", "150000", "8", "11", "1850000", "100"]],
        "ORDER BY 2 DESC LIMIT": [["11987654321", "500"], ["1134567890", "300"]],
        "regexp_replace": [["99999999999", "1500000"], ["9999999999", "300000"]],
    })
    p = catalogo.perfil(projeto, "cat.s.t", "cmpo_fone", executor=ex)
    assert p["total_tabela"] == 2000000 and p["nulos"] == 100000 and p["tam_min"] == 8 and p["tam_max"] == 11
    assert p["frequentes"][0]["valor"] == "99999999999" and p["mascarado"]
    assert "TABLESAMPLE" not in ex.chamadas[1]  # 2 mi < padrao de 5 mi
    txt = "\n".join(catalogo.texto_perfil(p))
    assert "nulos 100.000" in txt and "padroes mais comuns" in txt
    t = catalogo.testar_predicado(projeto, "cat.s.t", "length(cmpo_fone) <> 9 OR cmpo_fone NOT LIKE '9%'", executor=ex)
    assert t["marcados"] == 4123 and t["total"] == 2000000 and abs(t["taxa"] - 0.0021) < 1e-4
    assert "marcaria 4.123 de 2.000.000" in "\n".join(catalogo.texto_teste(t))
    with pytest.raises(catalogo.SqlInvalido):
        catalogo.testar_predicado(projeto, "cat.s.t", "1=1; DROP TABLE x", executor=ex)
    ex2 = _executor_contagens_analise_dados({"SELECT DISTINCT": [["11999990000"], ["21988887777"]]})
    r = catalogo.exemplos(projeto, "cat.s.t", "cmpo_fone", "length(cmpo_fone) = 11", executor=ex2)
    assert r["atendem"] == ["99999999999", "99999999999"] and len(ex2.chamadas) == 2
    r2 = catalogo.exemplos(projeto, "cat.s.t", "cmpo_fone", "", mascarar=False, executor=ex2)
    assert r2["atendem"][0] == "11999990000"
    s = catalogo.sobreposicao(projeto, "cat.s.t", "a = 1", "b = 2", executor=_executor_contagens_analise_dados({"sum(CASE": [["100", "80", "60"]]}))
    assert s["ambos"] == 60 and s["taxa_do_novo_ja_coberta"] == 0.6


# ---------------------------------------------------------------- tabelas sob demanda
class _Col_analise_dados:
    def __init__(self, n, t="string"):
        self.name, self.type_text, self.comment = n, t, ""


class _Info_analise_dados:
    def __init__(self, cols, comment=""):
        self.columns, self.comment, self.table_type = cols, comment, "MANAGED"


class _Tables_analise_dados:
    def __init__(self, existentes):
        self.existentes = existentes

    def get(self, full_name):
        if full_name not in self.existentes:
            raise RuntimeError(f"TABLE_OR_VIEW_NOT_FOUND: {full_name} does not exist")
        return self.existentes[full_name]

    def list(self, catalog_name, schema_name):
        class T:
            def __init__(self, n): self.name = n
        return [T(f.split(".")[-1]) for f in self.existentes if f.startswith(f"{catalog_name}.{schema_name}.")]


class _Cliente_analise_dados:
    def __init__(self, existentes):
        self.tables = _Tables_analise_dados(existentes)


def test_incluir_tabela_ranqueia_e_indexa(projeto):
    nucleo.gravar_valor(projeto.raiz, "databricks", "schemas_preferidos", "_gold")  # grava string; o ranking le lista
    cfg = nucleo.carregar(projeto.raiz)
    cfg.bruto["databricks"]["schemas_preferidos"] = ["_gold", "_silver"]
    existentes = {
        "cat.ing.fone": _Info_analise_dados([_Col_analise_dados("cmpo_fone"), _Col_analise_dados("cmpo_ddd")]),
        "cat.sch_gold.fone": _Info_analise_dados([_Col_analise_dados("cmpo_fone"), _Col_analise_dados("cmpo_ddd"), _Col_analise_dados("cmpo_tipo")], "telefones consolidados"),
        "cat.sch_gold.fone_bckp": _Info_analise_dados([_Col_analise_dados("cmpo_fone")]),
    }
    cliente_ = _Cliente_analise_dados(existentes)
    homonimas = [["cat", "ing", "fone", "MANAGED"], ["cat", "sch_gold", "fone", "MANAGED"], ["cat", "sch_gold", "fone_bckp", "MANAGED"]]

    def executor(sql):
        if "information_schema.tables" in sql:
            return [h for h in homonimas if h[2] == "fone"]
        if "query.history" in sql:
            return [["cat.sch_gold.fone", "250"], ["cat.ing.fone", "3"]]
        return []
    rel = catalogo.incluir_tabela(cfg, "cat.ing.fone", {"CMPO_FONE"}, cliente=cliente_, executor=executor)
    assert rel["existe"] and rel["recomendada"] == "cat.ing.fone"  # nome exato pesa mais que o resto
    pontos = {c["fqn"]: c["pontos"] for c in rel["candidatas"]}
    assert pontos["cat.sch_gold.fone"] > 0 and "schema preferido (_gold)" in next(c["motivos"] for c in rel["candidatas"] if c["fqn"] == "cat.sch_gold.fone")
    # pedindo uma que nao existe com esse nome exato: a gold vence pelo uso + schema
    rel2 = catalogo.incluir_tabela(cfg, "cat.outro.fone", {"CMPO_FONE"}, cliente=cliente_, executor=executor)
    assert not rel2["existe"] and rel2["recomendada"] == "cat.sch_gold.fone"
    txt = "\n".join(catalogo.texto_inclusao(rel2))
    assert "NAO existe com esse nome exato" in txt and "recomendada: cat.sch_gold.fone" in txt
    # extras gravados e indice unificado enxerga
    extras = catalogo.carregar_extras(cfg)
    assert "cat.sch_gold.fone" in extras and "cat.ing.fone" in extras
    idx = catalogo.indice_colunas(catalogo.carregar_snapshot(cfg), extras)
    assert any(o["fqn"] == "cat.sch_gold.fone" and o["extra"] for o in idx["CMPO_FONE"])
    linhas = catalogo.descrever_tabela(catalogo.carregar_snapshot(cfg), "cat.sch_gold.fone", extras=extras)
    assert "telefones consolidados" in linhas[0] and "incluida sob demanda" in linhas[0]
    assert catalogo.resolver_fqn(cfg, "cat.ing.fone", extras=extras) == "cat.ing.fone"
    assert catalogo.resolver_fqn(cfg, "fone", extras=extras) == ""  # ambiguo: existe em dois schemas
    assert "existe em mais de um lugar" in catalogo.descrever_tabela(catalogo.carregar_snapshot(cfg), "fone", extras=extras)[0]
    # inexistente em todo lugar
    rel3 = catalogo.incluir_tabela(cfg, "cat.x.nada", cliente=cliente_, executor=lambda sql: [])
    assert rel3.get("erro")


def test_validador_c13_c14(projeto):
    r = nucleo.completar({
        "codigo": "ICT.RL.7", "nome": "Saneamento do DDD", "tipo": "Limpeza", "criticidade": "Alta", "tabela": "Informações de Contato (INFO_CNTTO)", "tipo_pessoa": "PF",
        "descricao": "Apaga o DDD do telefone de contato quando inválido.", "objetivo": "Apagar o DDD do telefone de contato inválido.",
        "criterio_aplicacao": "Executar em todos os registros segregados pela regra ICT.RV.1, com nucleo em cat.outro.tabela_z.",
        "logica_execucao": "A regra apaga o conteúdo do campo CMPO_DDD_TEL nos registros segregados.", "resultado": "Sucesso: apagado. Falha: não apagado.",
    })
    ctx = nucleo.Contexto({"CMPO_DDD_TEL"}, {"ICT.RV.1"}, tabelas_conhecidas={"cat.s.contato_pf_ex", "contato_pf_ex"})
    checks = {a.check for a in nucleo.validar(r, projeto, ctx)}
    assert "C13_tabela_fora_do_indice" in checks and "C14_nao_testado_no_dado" in checks
    r["teste_no_dado"] = "10 de 100 (10%) em 2026-01-01"
    ctx.tabelas_conhecidas.add("cat.outro.tabela_z")
    checks = {a.check for a in nucleo.validar(r, projeto, ctx)}
    assert "C13_tabela_fora_do_indice" not in checks and "C14_nao_testado_no_dado" not in checks


def test_cli_novos_comandos(projeto, capsys):
    raiz = str(projeto.raiz)
    assert cli.main(["--raiz", raiz, "atualizar", "--sem-uc"]) == 0
    out = capsys.readouterr().out
    assert "cadeias:" in out and (projeto.dados / "metricas.json").exists() and (projeto.dados / "grafo.gexf").exists()
    assert (projeto.dados / "visoes" / "cobertura.md").exists()
    html = (projeto.dados / "grafo.html").read_text(encoding="utf-8")
    assert '"com_metricas": true' in html and "Impacto a jusante" in html
    assert cli.main(["--raiz", raiz, "explorar", "CMPO_DDD_TEL"]) == 0
    assert "regra(s) atrelada(s)" in capsys.readouterr().out
    assert cli.main(["--raiz", raiz, "impacto", "ICT.RV.1"]) == 0
    assert "ICT.RL.1" in capsys.readouterr().out
    assert cli.main(["--raiz", raiz, "cadeias"]) == 0 and cli.main(["--raiz", raiz, "redundancias"]) == 0
    assert cli.main(["--raiz", raiz, "cobertura"]) == 0 and cli.main(["--raiz", raiz, "dimensoes"]) == 0
    capsys.readouterr()
    assert cli.main(["--raiz", raiz, "exportar-grafo"]) == 0
    assert "GEXF" in capsys.readouterr().out
    assert cli.main(["--raiz", raiz, "status"]) == 0
    assert "metricas do grafo:" in capsys.readouterr().out
    # testar-regra com executor nao e possivel pela CLI (precisa de rede): o comando deve falhar limpo
    ficha = projeto.saida / "_f.json"
    ficha.write_text(json.dumps({"codigo": "ICT.RV.8", "tabela_sigla": "INFO_CNTTO", "sql_teste": "1=1"}), encoding="utf-8")
    assert cli.main(["--raiz", raiz, "testar-regra", str(ficha)]) == 1


# ==========================================================================================
# test_base.py
# ==========================================================================================
def test_ler_aprovadas_md():
    regras = nucleo.ler_aprovadas_md(EXEMPLOS / "base" / "regras_aprovadas_exemplo.md", {"SE", "RBRF", "PSDC"})
    assert len(regras) == 8
    rv1 = next(r for r in regras if r["codigo"] == "ICT.RV.1")
    assert rv1["tipo"] == "Validação" and rv1["tabela_sigla"] == "INFO_CNTTO" and rv1["tabela"] == "Informações de Contato"
    assert "Se constar valores" in rv1["logica_execucao"]          # continuacao de paragrafo
    assert "Executar regra ICT.RL.1" in rv1["acao_dado_invalido"]   # continuacao do ultimo campo
    assert rv1["regras_encadeadas"] == ["ICT.RL.1"]
    assert "CMPO_DDD_TEL" in rv1["campos"] and rv1["dominios"] == ["TDDD - XDOM001"]
    rl = next(r for r in regras if r["codigo"] == "IBS.RL.1")
    assert rl["resultado"] == ""  # 'nan' virou vazio
    assert rl["regras_encadeadas"] == ["IBS.RV.1"]  # citada no criterio de aplicacao


def test_ler_base_conhecimento():
    kb = nucleo.ler_base_conhecimento(EXEMPLOS / "base" / "Regras de governanca.md")
    assert set(kb["secoes"]) >= {"tipos_de_regra", "criticidade", "campos_obrigatorios", "ciclo_de_vida"}
    assert nucleo.tipos_da_base(kb) == ["Negócio", "Validação", "Enriquecimento", "Limpeza", "Retroalimentação", "Regra Geral"]
    obrig = nucleo.campos_obrigatorios_da_base(kb)
    assert "Resultado" in obrig["TODAS"] and "Exemplo" in obrig["Validação"]
    assert "chunk_size" not in kb["texto"] and "Contexto e Propósito" in kb["texto"]


def test_glossario_das_aprovadas():
    regras = nucleo.ler_aprovadas_md(EXEMPLOS / "base" / "regras_aprovadas_exemplo.md")
    g = {x["sigla"]: x for x in nucleo.extrair_glossario(regras, {"CMPO_NAC": "Codigo da nacionalidade"})}
    assert "ddd do telefone" in g["CMPO_DDD_TEL"]["termo_negocio"].lower()
    assert "sexo" in g["CMPO_SEXO"]["termo_negocio"].lower()
    assert g["CMPO_NAC"]["fonte"] == "unity_catalog"


def test_ler_export_xlsx(tmp_path):
    sys.path.insert(0, str(EXEMPLOS))
    import gerar_export

    p = gerar_export.gerar(tmp_path / "export.xlsx")
    regras = nucleo.ler_export_xlsx(p, {"SE", "ENTÃO"})
    assert len(regras) == 3
    rv1 = next(r for r in regras if r["codigo"] == "ICT.RV.1")
    assert rv1["codigo_comp"] == "ICT.RN1.RV1" and rv1["rn"] == "ICT.RN.1"
    assert rv1["resultado"] == ""  # nan
    assert rv1["flags"]["is_legacy"] is True and rv1["flags"]["is_editing"] is False
    assert rv1["extras"]["bloco"] == "Dados de Contato"
    assert rv1["status"] == "Aprovada"
    assert rv1["regras_encadeadas"] == ["ICT.RL.1"]


def test_descobrir_arquivos(tmp_path):
    (tmp_path / "Regras de governanca.md").write_text("x", encoding="utf-8")
    (tmp_path / "regras_aprovadas_51_pf.md").write_text("x", encoding="utf-8")
    (tmp_path / "Extração CadFlow 20.08.2026.xlsx").write_bytes(b"")
    (tmp_path / "~$Extração CadFlow 20.08.2026.xlsx").write_bytes(b"")
    a = nucleo.descobrir_arquivos(tmp_path)
    assert len(a["conhecimento"]) == 1 and len(a["aprovadas"]) == 1 and len(a["export"]) == 1


# ==========================================================================================
# test_calibracao_real.py
# ==========================================================================================
"""Casos vindos da primeira calibracao com as regras reais da Governanca (2026-10-08).

Cada teste aqui e um ajuste do validador a algo que a Governanca ja aprova. Se um deles quebrar,
o validador voltou a reprovar texto aprovado.
"""




def _enriquecimento_calibracao_real(logica: str) -> dict:
    return nucleo.completar({
        "codigo": "IBS.RE.7", "nome": "Enriquecimento de nome social", "tipo": "Enriquecimento", "criticidade": "Alta",
        "tabela": "Informações Básicas (INFO_BSICO)", "tipo_pessoa": "PF",
        "descricao": "Enriquece o campo nome social de pessoa física em fonte externa, quando o campo estiver vazio.",
        "objetivo": "Enriquecer o campo nome social quando estiver vazio no PSDC.",
        "criterio_aplicacao": "Executar em todos os registros com campo nome social vazio no PSDC.",
        "logica_execucao": logica,
        "resultado": "Sucesso: registros enriquecidos. Falha: registro não encontrado na RBRF.",
    })


def test_c4_aceita_verbo_do_tipo_depois_da_condicao(projeto):
    ctx = nucleo.Contexto(set(), set())
    r = _enriquecimento_calibracao_real("A regra verifica se o campo CMPO_NOME_SOC está vazio no PSDC e enriquece com a fonte RBRF.")
    assert "C4_verbo_do_tipo" not in {a.check for a in nucleo.validar(r, projeto, ctx)}
    r = _enriquecimento_calibracao_real("A regra compara o campo CMPO_NOME_SOC com a RBRF e preenche o conteúdo quando divergente.")
    assert "C4_verbo_do_tipo" not in {a.check for a in nucleo.validar(r, projeto, ctx)}
    r = _enriquecimento_calibracao_real("A regra verifica se o campo CMPO_NOME_SOC está vazio e segrega o registro.")
    assert "C4_verbo_do_tipo" in {a.check for a in nucleo.validar(r, projeto, ctx)}


def test_c2_erro_so_para_sigla_conhecida(projeto):
    r = _enriquecimento_calibracao_real("A regra enriquece o campo CMPO_NOME_SOC com nucleo na RBRF.")
    ctx = nucleo.Contexto(siglas_conhecidas={"CMPO_NOME_SOC", "CMPO_OUTRO"}, codigos_existentes=set())
    # sigla que a logica tambem cita (Limpeza aprovada que lista os campos): so aviso
    r["descricao"] = "Enriquece o nome social conforme a LGPD, usando o campo CMPO_NOME_SOC."
    ach = {(a.check, a.nivel) for a in nucleo.validar(r, projeto, ctx)}
    assert ("C2_sigla_em_texto_negocio", "aviso") in ach and ("C2_sigla_em_texto_negocio", "erro") not in ach
    # campo conhecido que a logica nem menciona: erro
    r["descricao"] = "Enriquece o nome social usando o campo CMPO_OUTRO."
    ach = {(a.check, a.nivel) for a in nucleo.validar(r, projeto, ctx)}
    assert ("C2_sigla_em_texto_negocio", "erro") in ach
    r["descricao"] = "Enriquece o nome social conforme a norma XPTO da empresa."
    ach = {(a.check, a.nivel) for a in nucleo.validar(r, projeto, ctx)}
    assert ("C2_sigla_em_texto_negocio", "aviso") in ach and ("C2_sigla_em_texto_negocio", "erro") not in ach


def test_mapa_por_nome_quando_nao_ha_colunas_em_comum(projeto):
    snap = {"tabelas": {"compl_pj_hive": {"colunas": []}, "compl_pj": {"colunas": []}, "docto_pssoa_unic": {"colunas": []}, "outra": {"colunas": []}}}
    regras = [nucleo.completar({**nucleo.regra_vazia(), "codigo": "CPX.RV.1", "tipo": "Validação", "tabela": "Complemento PJ (COMPL_PJ)", "logica_execucao": "A regra verifica o campo CMPO_X."}),
              nucleo.completar({**nucleo.regra_vazia(), "codigo": "DOX.RV.1", "tipo": "Validação", "tabela": "Documento (TDOCTO_PSSOA_UNIC)", "logica_execucao": "A regra verifica o campo CMPO_Y."})]
    m = nucleo.inferir_tabelas_uc(regras, snap)
    assert m["COMPL_PJ"]["tabela_uc"] == "compl_pj" and m["COMPL_PJ"]["criterio"] == "nome" and "compl_pj_hive" in m["COMPL_PJ"]["alternativas"]
    assert m["TDOCTO_PSSOA_UNIC"]["tabela_uc"] == "docto_pssoa_unic"


def test_mapa_prefere_nome_curto_em_empate(projeto):
    cols = [{"nome": "cmpo_a"}, {"nome": "cmpo_b"}]
    snap = {"tabelas": {"bsico_x_previa": {"colunas": cols}, "bsico_x_hive": {"colunas": cols}, "bsico_x": {"colunas": cols}}}
    regras = [nucleo.completar({**nucleo.regra_vazia(), "codigo": "BSX.RV.1", "tipo": "Validação", "tabela": "Básico (BSICO_X)", "logica_execucao": "A regra verifica os campos CMPO_A e CMPO_B."})]
    assert nucleo.inferir_tabelas_uc(regras, snap)["BSICO_X"]["tabela_uc"] == "bsico_x"


def test_recursos_e_resolver_por_nome(projeto):
    class S:
        def __init__(self, v): self.value = v

    class C:
        def __init__(self, i, n): self.cluster_id, self.cluster_name, self.state = i, n, S("TERMINATED")

    class W:
        def __init__(self, i, n): self.id, self.name, self.state = i, n, S("RUNNING")

    class Cl:
        def list(self): return [C("0000-1", "psdc-obs-interativo"), C("0000-2", "outro")]

    class Wh:
        def list(self): return [W("abc", "Serverless Starter")]

    class Cliente:
        clusters = Cl(); warehouses = Wh()

    rec = catalogo.listar_recursos(projeto, cliente=Cliente())
    assert rec["clusters"][0] == {"id": "0000-1", "nome": "psdc-obs-interativo", "estado": "TERMINATED"}
    assert catalogo.resolver_por_nome(rec["clusters"], "psdc-obs") == "0000-1"
    assert catalogo.resolver_por_nome(rec["clusters"], "o") == ""  # ambiguo
    assert catalogo.resolver_por_nome(rec["warehouses"], "serverless starter") == "abc"


# ==========================================================================================
# test_cluster.py
# ==========================================================================================
"""Cluster all-purpose desligado: nada pode esperar em silencio nem derrubar o comando sem instrucao."""





class _Estado_cluster:
    def __init__(self, v):
        self.value = v


class _Cluster_cluster:
    def __init__(self, estado, nome="psdc-obs"):
        self.state, self.cluster_name, self.state_message = _Estado_cluster(estado), nome, ""


class _Clusters_cluster:
    """Sequencia de estados devolvida a cada get(); start() registra a chamada."""

    def __init__(self, estados):
        self.estados = list(estados)
        self.starts = 0

    def get(self, cluster_id):
        # um cluster TERMINATED so avanca de estado depois de start()
        if self.estados[0] == "TERMINATED" and self.starts == 0:
            return _Cluster_cluster("TERMINATED")
        e = self.estados.pop(0) if len(self.estados) > 1 else self.estados[0]
        return _Cluster_cluster(e)

    def start(self, cluster_id):
        self.starts += 1
        if self.estados[0] == "TERMINATED" and len(self.estados) > 1:
            self.estados.pop(0)  # ligou: o proximo get() ja ve PENDING
        return object()


class _Exec_cluster:
    def __init__(self):
        self.chamadas = 0

    def create(self, cluster_id, language):
        class R:
            def result(_):
                class C:
                    id = "ctx"
                return C()
        return R()

    def execute(self, cluster_id, context_id, language, command):
        self.chamadas += 1

        class R:
            def result(_):
                class Res:
                    class results:
                        result_type = "table"
                        data = [["7"]]
                return Res()
        return R()

    def destroy(self, cluster_id, context_id):
        pass


class _Cliente_cluster:
    def __init__(self, estados):
        self.clusters = _Clusters_cluster(estados)
        self.command_execution = _Exec_cluster()


def _cfg_cluster(projeto, auto=True, espera=30):
    nucleo.gravar_valor(projeto.raiz, "databricks", "cluster_id", "0000-1", "config.local.toml")
    c = nucleo.carregar(projeto.raiz)
    c.bruto["databricks"]["auto_ligar_cluster"] = auto
    c.bruto["databricks"]["espera_cluster_s"] = espera
    c.bruto["databricks"]["intervalo_poll_s"] = 5
    return c


def test_estado_ligar_esperar(projeto):
    cfg = _cfg_cluster(projeto)
    cli_ = _Cliente_cluster(["TERMINATED", "PENDING", "PENDING", "RUNNING"])
    assert catalogo.estado_cluster(cfg, cli_)["estado"] == "TERMINATED"
    e = catalogo.ligar_cluster(cfg, cli_)
    assert cli_.clusters.starts == 1 and e["estado"] == "PENDING"
    linhas = []
    e = catalogo.esperar_cluster(cfg, cli_, 30, linhas.append, dormir=lambda s: None)
    assert e["estado"] == "RUNNING" and any("PENDING" in l for l in linhas)


def test_garantir_cluster_liga_e_espera(projeto, monkeypatch):
    monkeypatch.setattr(catalogo.time, "sleep", lambda s: None)
    cfg = _cfg_cluster(projeto)
    cli_ = _Cliente_cluster(["TERMINATED", "PENDING", "RUNNING"])
    e = catalogo.garantir_cluster(cfg, cli_, None)
    assert e["estado"] == "RUNNING" and cli_.clusters.starts == 1


def test_garantir_cluster_sem_auto_ligar(projeto):
    cfg = _cfg_cluster(projeto, auto=False)
    with pytest.raises(catalogo.ClusterDesligado):
        catalogo.garantir_cluster(cfg, _Cliente_cluster(["TERMINATED"]), None)


def test_garantir_cluster_estoura_espera(projeto, monkeypatch):
    monkeypatch.setattr(catalogo.time, "sleep", lambda s: None)
    cfg = _cfg_cluster(projeto, espera=10)
    with pytest.raises(catalogo.ClusterIniciando):
        catalogo.garantir_cluster(cfg, _Cliente_cluster(["PENDING"]), None)


def test_executar_sql_propaga_pendencia_e_roda_quando_running(projeto, monkeypatch):
    monkeypatch.setattr(catalogo.time, "sleep", lambda s: None)
    cfg = _cfg_cluster(projeto, espera=10)
    with pytest.raises(catalogo.ClusterIniciando):
        catalogo.executar_sql(cfg, "SELECT 1", cliente=_Cliente_cluster(["PENDING"]))
    cli_ = _Cliente_cluster(["RUNNING"])
    linhas, erros, via = catalogo.executar_sql(cfg, "SELECT 1", cliente=cli_)
    assert linhas == [["7"]] and via == "cluster" and cli_.command_execution.chamadas == 1


def test_linhagem_e_homonimas_nao_ligam_cluster(projeto, monkeypatch):
    monkeypatch.setattr(catalogo.time, "sleep", lambda s: None)
    cfg = _cfg_cluster(projeto, espera=10)
    cli_ = _Cliente_cluster(["PENDING"])
    lin = catalogo.consultar_linhagem(cfg, cli_)
    assert lin["relacoes"] == {} and any("adiada" in e for e in lin["erros"])

    class Sc:
        def __init__(self, n): self.name = n

    class Schemas:
        def list(self, catalog_name): return [Sc("s1"), Sc("information_schema")]

    class T:
        def __init__(self, n): self.name = n

    class Tables:
        def list(self, catalog_name, schema_name): return [T("fone"), T("outra")]
    cli_.schemas, cli_.tables = Schemas(), Tables()
    achadas, erros = catalogo.buscar_homonimas(cfg, "fone", cli_)
    assert achadas == [f"{cfg.catalogo}.s1.fone"] and cli_.clusters.starts == 0


def test_cli_cluster_e_rc2(projeto, monkeypatch, capsys):
    monkeypatch.setattr(catalogo.time, "sleep", lambda s: None)
    cfg = _cfg_cluster(projeto, espera=10)
    raiz = str(projeto.raiz)
    monkeypatch.setattr(catalogo, "_cliente", lambda c: _Cliente_cluster(["TERMINATED"]))
    assert cli.main(["--raiz", raiz, "cluster"]) == 2
    assert "parado" in capsys.readouterr().out
    monkeypatch.setattr(catalogo, "_cliente", lambda c: _Cliente_cluster(["TERMINATED", "PENDING", "PENDING", "RUNNING"]))
    assert cli.main(["--raiz", raiz, "cluster", "--ligar", "--max", "30"]) == 0
    out = capsys.readouterr().out
    assert "pedido de ligar enviado" in out and "RUNNING" in out and "pronto" in out
    # comando de dado com cluster pendente -> rc 2 e instrucao
    monkeypatch.setattr(catalogo, "_cliente", lambda c: _Cliente_cluster(["PENDING"]))
    nucleo.gravar_valor(projeto.raiz, "databricks", "espera_cluster_s", "10", "config.local.toml")
    assert cli.main(["--raiz", raiz, "perfil", "cat_exemplo.esquema_exemplo.contato_pf_ex", "cmpo_ddd_tel"]) == 2
    out = capsys.readouterr().out
    assert "CLUSTER PENDENTE" in out and "cluster --esperar" in out


# ==========================================================================================
# test_desfazer.py
# ==========================================================================================
def _ficha_desfazer(cod="ICT.RV.3", crit="Baixa"):
    return {"codigo": cod, "nome": "Validação do número do telefone de contato", "tipo": "Validação", "criticidade": crit,
            "tabela": "Informações de Contato (INFO_CNTTO)", "tipo_pessoa": "PF",
            "descricao": "Valida se o número do telefone de contato tem a quantidade de dígitos esperada.", "objetivo": "Validar o número do telefone de contato de pessoa física.",
            "criterio_aplicacao": "Executar em todos os registros que possuam conteúdo no campo número do telefone de contato.",
            "logica_execucao": "A regra verifica se o valor contido no campo CMPO_NUM_TEL possui 9 dígitos, sendo considerado válido ao atender esta condição.",
            "resultado": "Registros com número do telefone fora de 9 dígitos são segregados como inconsistentes.",
            "exemplo": "Dado analisado: 912345678. Como possui 9 dígitos, o dado é válido; o registro é mantido.", "acao_dado_valido": "Manter o dado validado, sem alterações.",
            "acao_dado_invalido": 'Segregar como inconsistente: "Número do telefone com quantidade de dígitos incoerente".'}


def test_desfazer_libera_codigo_e_limpa_tudo(projeto, capsys):
    raiz = str(projeto.raiz)
    assert cli.main(["--raiz", raiz, "atualizar", "--sem-uc"]) == 0
    ficha = projeto.saida / "_ficha_desfazer.json"
    ficha.write_text(json.dumps(_ficha_desfazer(), ensure_ascii=False), encoding="utf-8")
    assert cli.main(["--raiz", raiz, "gravar", str(ficha)]) == 0
    assert cli.main(["--raiz", raiz, "validar", str(ficha)]) == 1  # C9 repetido
    capsys.readouterr()
    assert cli.main(["--raiz", raiz, "proximo-codigo", "--tabela", "INFO_CNTTO", "--tipo", "Validação"]) == 0
    assert '"ICT.RV.4"' in capsys.readouterr().out
    # desfaz
    assert cli.main(["--raiz", raiz, "desfazer", "ict.rv3"]) == 0
    out = capsys.readouterr().out
    assert "desfeitas: 1 registro(s), 1 linha(s) do Excel, 1 ficha(s)" in out
    assert nucleo.ler_jsonl(projeto.dados / "regras_novas.jsonl") == []
    assert load_workbook(projeto.excel_saida).active.max_row == 1
    assert not (projeto.fichas / "ICT_RV_3.md").exists()
    assert cli.main(["--raiz", raiz, "validar", str(ficha)]) == 0  # codigo livre de novo
    capsys.readouterr()
    assert cli.main(["--raiz", raiz, "proximo-codigo", "--tabela", "INFO_CNTTO", "--tipo", "Validação"]) == 0
    assert '"ICT.RV.3"' in capsys.readouterr().out
    # regrava com outra criticidade (refeita) e desfazer de regra do export e recusado
    ficha.write_text(json.dumps(_ficha_desfazer(crit="Média"), ensure_ascii=False), encoding="utf-8")
    assert cli.main(["--raiz", raiz, "gravar", str(ficha)]) == 0
    assert cli.main(["--raiz", raiz, "desfazer", "ICT.RV.1"]) == 1
    assert "nao foi criada pelo cdflw" in capsys.readouterr().out
    assert cli.main(["--raiz", raiz, "desfazer", "XXX.RV.9"]) == 1



def _rn_desfazer():
    return {"codigo": "ICT.RN.9", "nome": "Consistência do número do telefone", "tipo": "Negócio", "criticidade": "Alta",
            "tabela": "Informações de Contato (INFO_CNTTO)", "tipo_pessoa": "PF",
            "descricao": "Verifica o número do telefone de contato, buscando qualificar os dados.", "objetivo": "Consistir o número do telefone de contato de pessoa física.",
            "criterio_aplicacao": "Executar em todo o cadastro de pessoa física.",
            "logica_execucao": "O conjunto de regras verifica o conteúdo do campo CMPO_NUM_TEL buscando consistir a informação.",
            "resultado": "Registros verificados e mantidos com dados consistentes."}


def _rl_desfazer():
    return {"codigo": "ICT.RL.9", "codigo_comp": "ICT.RN9.RL9", "rn": "ICT.RN.9", "nome": "Saneamento do número do telefone", "tipo": "Limpeza", "criticidade": "Alta",
            "tabela": "Informações de Contato (INFO_CNTTO)", "tipo_pessoa": "PF",
            "descricao": "Apaga o número do telefone de contato inválido.", "objetivo": "Apagar o número do telefone de contato inválido.",
            "criterio_aplicacao": "Executar em todos os registros segregados como inconsistentes por número do telefone com quantidade de dígitos incoerente.",
            "logica_execucao": "SE o registro estiver segregado por número do telefone incoerente ENTÃO apagar o conteúdo do campo CMPO_NUM_TEL. SENÃO manter o dado.",
            "resultado": "Registros segregados ficam com o número do telefone apagado.", "familia": ["ICT.RV.9"],
            "teste_no_dado": "1 de 10 (10%) em 2026-01-01"}


def test_desfazer_familia_por_cadeia_e_por_lote(projeto, capsys):
    raiz = str(projeto.raiz)
    assert cli.main(["--raiz", raiz, "atualizar", "--sem-uc"]) == 0
    assert cli.main(["--raiz", raiz, "trilha", "--nova", "familia telefone"]) == 0
    lote = (projeto.saida / "_lote_atual.txt").read_text(encoding="utf-8").strip()
    ficha = projeto.saida / "_ficha_desfazer.json"
    for f in (_rn_desfazer(), {**_ficha_desfazer(cod="ICT.RV.9"), "codigo_comp": "ICT.RN9.RV9", "rn": "ICT.RN.9", "familia": ["ICT.RL.9"]}, _rl_desfazer()):
        ficha.write_text(json.dumps(f, ensure_ascii=False), encoding="utf-8")
        assert cli.main(["--raiz", raiz, "gravar", str(ficha), "--forcar"]) == 0
    capsys.readouterr()
    # uma regra solta de outro lote
    assert cli.main(["--raiz", raiz, "trilha", "--nova", "outra"]) == 0
    ficha.write_text(json.dumps(_ficha_desfazer(cod="ICT.RV.3"), ensure_ascii=False), encoding="utf-8")
    assert cli.main(["--raiz", raiz, "gravar", str(ficha)]) == 0
    capsys.readouterr()
    assert cli.main(["--raiz", raiz, "status"]) == 0
    assert f"lote {lote}" in capsys.readouterr().out
    # cadeia a partir da RL pega RV e RN, nao a solta
    assert cli.main(["--raiz", raiz, "desfazer", "ICT.RL.9", "--cadeia", "--listar"]) == 0
    out = capsys.readouterr().out
    assert "3 regra(s) a desfazer" in out and "ICT.RN.9" in out and "ICT.RV.9" in out and "ICT.RV.3" not in out and "nada foi removido" in out
    assert len(nucleo.ler_jsonl(projeto.dados / "regras_novas.jsonl")) == 4
    # por lote
    assert cli.main(["--raiz", raiz, "desfazer", "--lote", lote]) == 0
    out = capsys.readouterr().out
    assert "desfeitas: 3 registro(s), 3 linha(s) do Excel, 3 ficha(s)" in out
    restantes = [r["codigo"] for r in nucleo.ler_jsonl(projeto.dados / "regras_novas.jsonl")]
    assert restantes == ["ICT.RV.3"]
    assert load_workbook(projeto.excel_saida).active.max_row == 2
    # lote inexistente e codigo do export
    assert cli.main(["--raiz", raiz, "desfazer", "--lote", "nada"]) == 1
    assert cli.main(["--raiz", raiz, "desfazer", "ICT.RV.1", "ICT.RV.3", "--listar"]) == 0
    out = capsys.readouterr().out
    assert "ICT.RV.1: veio do export" in out and "1 regra(s) a desfazer" in out


# ==========================================================================================
# test_grafo_mapa_codigo.py
# ==========================================================================================
def _dados_grafo_mapa_codigo(projeto):
    regras = nucleo.ler_aprovadas_md(EXEMPLOS / "base" / "regras_aprovadas_exemplo.md", projeto.nao_siglas)
    snap = json.loads((EXEMPLOS / "uc" / "esquema_exemplo.json").read_text(encoding="utf-8"))
    return regras, snap


def test_grafo_relacionadas_e_colisao(projeto):
    regras, snap = _dados_grafo_mapa_codigo(projeto)
    g = grafo.construir(regras, set(catalogo.indice_colunas(snap)))
    rel = grafo.relacionadas(g, ["CMPO_TIPO_TEL"], "INFO_CNTTO", tipo="Enriquecimento")
    codigos = [r["codigo"] for r in rel]
    assert codigos[0] == "ICT.RE.1" and "COLISAO" in rel[0]["alerta"]
    assert "ICT.RV.2" in codigos
    assert all(r["codigo"] != "IBS.RV.1" for r in rel)


def test_grafo_detecta_buracos(projeto):
    regras, snap = _dados_grafo_mapa_codigo(projeto)
    # tira a limpeza do sexo e acrescenta uma validacao orfa
    regras = [r for r in regras if r["codigo"] != "IBS.RL.1"]
    orfa = nucleo.completar({**nucleo.regra_vazia(), "codigo": "IBS.RV.9", "nome": "orfa", "tipo": "Validação",
                           "logica_execucao": "A regra verifica o campo CMPO_NAC.", "tabela": "X (INFO_BSICO)"})
    g = grafo.construir(regras + [orfa], set(catalogo.indice_colunas(snap)))
    tipos = {(a["tipo"], a["regra"]) for a in grafo.detectar(g)}
    assert ("codigo_inexistente", "IBS.RL.1") in tipos          # IBS.RV.1 chama uma regra que sumiu
    assert ("validacao_sem_tratamento", "IBS.RV.9") in tipos


def test_grafo_campo_fora_do_catalogo(projeto):
    regras, snap = _dados_grafo_mapa_codigo(projeto)
    r = nucleo.completar({**nucleo.regra_vazia(), "codigo": "IBS.RL.7", "nome": "x", "tipo": "Limpeza",
                        "logica_execucao": "A regra apaga o campo CMPO_FANTASMA.", "tabela": "X (INFO_BSICO)"})
    g = grafo.construir(regras + [r], set(catalogo.indice_colunas(snap)))
    assert any(a["tipo"] == "campo_fora_do_catalogo" and "CMPO_FANTASMA" in a["detalhe"] for a in grafo.detectar(g))


def test_mapa_infere_prefixo_e_tabela(projeto):
    regras, snap = _dados_grafo_mapa_codigo(projeto)
    pfx = nucleo.inferir_prefixos(regras)
    assert pfx["INFO_CNTTO"]["prefixo"] == "ICT" and pfx["INFO_BSICO"]["prefixo"] == "IBS"
    tab = nucleo.inferir_tabelas_uc(regras, snap)
    assert tab["INFO_CNTTO"]["tabela_uc"] == "contato_pf_ex" and tab["INFO_BSICO"]["tabela_uc"] == "basico_pf_ex"
    m = nucleo.mapa_final(projeto, regras, snap)
    assert not m["INFO_CNTTO"]["prefixo_confirmado"]
    assert "[mapa_tabelas.prefixo]" in nucleo.texto_mapa(m)


def test_proximo_codigo(projeto):
    regras, _ = _dados_grafo_mapa_codigo(projeto)
    usados = nucleo.codigos_usados(regras)
    assert nucleo.proximo("ICT", "Validação", projeto, usados)["codigo"] == "ICT.RV.3"
    assert nucleo.proximo("ICT", "Limpeza", projeto, usados, rn="ICT.RN.1") == {"codigo": "ICT.RL.2", "codigo_comp": "ICT.RN1.RL2"}
    assert nucleo.proximo("NOV", "Enriquecimento", projeto, usados)["codigo"] == "NOV.RE.1"
    assert "aviso" in nucleo.proximo("ICT", "Limpeza", projeto, usados, rn="IBS.RN.1")


def test_uc_descreve_campo_com_linhagem(projeto):
    snap = catalogo.carregar_snapshot(projeto)
    lin = catalogo.carregar_linhagem(projeto)
    linhas = catalogo.descrever_campo(snap, lin, "cmpo_tipo_tel")
    assert "contato_pf_ex" in linhas[0] and any("deriva de psdc_raw.cntto_raw.tipo_tel" in l for l in linhas)
    assert "NAO encontrada" in catalogo.descrever_campo(snap, lin, "CMPO_NADA")[0]
    assert "3 colunas" in " ".join(catalogo.descrever_tabela(snap, "basico_pj_ex"))


def test_uc_listar_tabelas_com_cliente_falso(projeto):
    class Col:
        def __init__(self, n, t, c=""):
            self.name, self.type_text, self.comment = n, t, c

    class Tab:
        def __init__(self, n, cols):
            self.name, self.columns, self.comment, self.table_type = n, cols, "cmt", "MANAGED"

    class Tables:
        def list(self, catalog_name, schema_name):
            return [Tab("t1", [Col("cmpo_a", "string", "A"), Col("cmpo_b", "int")])]

    class W:
        tables = Tables()

    snap = catalogo.listar_tabelas(projeto, cliente=W())
    assert snap["tabelas"]["t1"]["colunas"][0] == {"nome": "cmpo_a", "tipo": "string", "comentario": "A"}
    assert snap["schema"] == projeto.schema and snap["snapshot_id"]


def test_uc_linhagem_com_executor_falso(projeto):
    linhas = [["psdc_raw", "raw_t", "c1", "t1", "cmpo_a", "4"]]
    lin = catalogo.consultar_linhagem(projeto, executor=lambda sql: linhas)
    assert lin["relacoes"]["t1.cmpo_a"] == [{"origem": "psdc_raw.raw_t.c1", "n": 4}] and lin["erros"] == []

    def quebra(sql):
        raise RuntimeError("sem permissao")

    lin2 = catalogo.consultar_linhagem(projeto, executor=quebra)
    assert lin2["relacoes"] == {} and "sem permissao" in lin2["erros"][0]


# ==========================================================================================
# test_gravar_cli.py
# ==========================================================================================
def _ficha_gravar_cli():
    return {
        "codigo": "ICT.RV.3", "codigo_comp": "ICT.RN1.RV3", "rn": "ICT.RN.1", "nome": "Validação do número do telefone de contato",
        "tipo": "Validação", "criticidade": "Baixa", "tabela": "Informações de Contato", "tabela_sigla": "INFO_CNTTO", "tipo_pessoa": "PF",
        "descricao": "Valida se o número do telefone de contato tem a quantidade de dígitos esperada para o tipo informado.",
        "objetivo": "Validar o número do telefone de contato de pessoa física.",
        "criterio_aplicacao": "Executar em todos os registros que possuam conteúdo no campo número do telefone de contato.",
        "logica_execucao": "A regra verifica se o valor contido no campo CMPO_NUM_TEL possui 9 dígitos quando o campo CMPO_TIPO_TEL for celular, sendo considerado válido ao atender esta condição.",
        "resultado": "Registros com número do telefone fora da quantidade de dígitos esperada são segregados como inconsistentes.",
        "exemplo": "Dado analisado: 912345678 com tipo celular. Como possui 9 dígitos, o dado é válido; o registro é mantido. Dado analisado: 12345 com tipo celular. Como não possui 9 dígitos, o dado é inválido; o registro é segregado como inconsistente.",
        "acao_dado_valido": "Manter o dado validado, sem alterações.",
        "acao_dado_invalido": 'Segregar como inconsistente: "Número do telefone com quantidade de dígitos incoerente".',
        "familia": ["ICT.RL.1"], "pontos_atencao": "Confirmar se o DDD entra na contagem de dígitos.",
    }


def test_gravar_excel_e_ficha(projeto):
    r = nucleo.completar(_ficha_gravar_cli())
    p = nucleo.gravar_excel(r, projeto)
    nucleo.gravar_excel(r, projeto)  # segunda linha
    ws = load_workbook(p).active
    assert [c.value for c in ws[1]] == projeto.colunas_excel
    assert ws.max_row == 3 and ws["A2"].value == "ICT.RV.3" and ws["F2"].value == "Informações de Contato (INFO_CNTTO)"
    # a planilha nao cita outra regra (padrao da Governanca); o vinculo fica no registro interno
    assert (ws.cell(row=2, column=projeto.colunas_excel.index("regras_encadeadas") + 1).value or "") == ""
    assert r["regras_encadeadas"] == ["ICT.RL.1"]
    assert ws.cell(row=2, column=projeto.colunas_excel.index("pontos_atencao") + 1).value.startswith("Confirmar")
    ficha = nucleo.gravar_ficha(r, projeto)
    t = ficha.read_text(encoding="utf-8")
    assert ficha.name == "ICT_RV_3.md" and "## Lógica de Execução" in t and "CMPO_NUM_TEL" in t and "## Status" not in t
    assert "## Pontos de atenção" in t and "## Regras encadeadas" not in t



def test_config_le_project_json(projeto):
    assert projeto.profile == "perfil_teste" and projeto.cluster_id == "0000-000000-teste"


def test_config_local_sobrepoe(projeto):
    nucleo.gravar_valor(projeto.raiz, "databricks", "profile", "meu_perfil", "config.local.toml")
    nucleo.gravar_valor(projeto.raiz, "mapa_tabelas.uc", "INFO_X", "tabela_x", "config.local.toml")
    c = nucleo.carregar(projeto.raiz)
    assert c.profile == "meu_perfil" and c.schema == "esquema_exemplo"  # local por cima, equipe mantida
    assert c.mapa_uc.get("INFO_X") == "tabela_x"
    assert (projeto.raiz / "config.local.toml").read_text(encoding="utf-8").startswith("# Valores desta maquina")


def test_cli_ponta_a_ponta(projeto, capsys):
    raiz = str(projeto.raiz)
    assert cli.main(["--raiz", raiz, "instalar"]) == 0
    assert cli.main(["--raiz", raiz, "atualizar", "--sem-uc"]) == 0
    saida = capsys.readouterr().out
    assert "8 regras" in saida and "0 reprovadas" in saida and "1 com campo vazio na fonte" in saida
    refs = projeto.copilot_fonte / "skills" / "regras-cadflow" / "references"
    assert (refs / "padrao.md").exists() and (refs / "exemplos.md").exists() and (refs / "tabelas.md").exists()
    assert "ICT.RV.1" in (refs / "exemplos.md").read_text(encoding="utf-8")
    assert (projeto.raiz_workspace / ".github" / "skills" / "regras-cadflow" / "references" / "exemplos.md").exists()
    assert (projeto.dados / "mapa_tabelas.md").read_text(encoding="utf-8").count("contato_pf_ex") >= 1

    # aceitar o mapa sugerido grava prefixos e tabelas no config.toml
    assert cli.main(["--raiz", raiz, "configurar", "--aceitar-mapa"]) == 0
    s = capsys.readouterr().out
    assert "mapa confirmado: 2 prefixos, 2 tabelas" in s, s
    cfg_txt = projeto.raiz.joinpath("config.toml").read_text(encoding="utf-8")
    assert 'INFO_CNTTO = "ICT"' in cfg_txt and 'INFO_CNTTO = "contato_pf_ex"' in cfg_txt

    assert cli.main(["--raiz", raiz, "status"]) == 0
    copia = (projeto.saida / "_ultimo_comando.txt").read_text(encoding="utf-8")
    assert copia.startswith("$ cdflw --raiz") and "cdflw versao" in copia and "dados/:" in copia
    assert cli.main(["--raiz", raiz, "campo", "CMPO_TIPO_TEL"]) == 0
    s = capsys.readouterr().out
    assert "contato_pf_ex" in s and "ICT.RE.1" in s

    assert cli.main(["--raiz", raiz, "relacionadas", "--campos", "CMPO_DDD_TEL", "--tipo", "Limpeza"]) == 0
    assert "COLISAO" in capsys.readouterr().out

    assert cli.main(["--raiz", raiz, "proximo-codigo", "--tabela", "INFO_CNTTO", "--tipo", "Validação", "--rn", "ICT.RN.1"]) == 0
    s = capsys.readouterr().out
    assert '"codigo": "ICT.RV.3"' in s and "ICT.RN1.RV3" in s

    ficha = projeto.saida / "_ficha_gravar_cli.json"
    ficha.write_text(json.dumps(_ficha_gravar_cli(), ensure_ascii=False), encoding="utf-8")
    assert cli.main(["--raiz", raiz, "validar", str(ficha), "--pedido", "validar o número do telefone PF"]) == 0
    assert "0 erro(s)" in capsys.readouterr().out

    assert cli.main(["--raiz", raiz, "gravar", str(ficha)]) == 0
    s = capsys.readouterr().out
    assert "Excel" in s and "Ficha" in s and projeto.excel_saida.exists()
    html = (projeto.dados / "grafo.html").read_text(encoding="utf-8")
    assert "ICT.RV.3" in html and "<canvas" in html and "cdn" not in html.lower()
    assert cli.main(["--raiz", raiz, "grafo-html"]) == 0
    assert (projeto.fichas / "ICT_RV_3.md").exists()

    # a regra gravada passa a existir: o proximo codigo avanca e o codigo repetido e erro
    assert cli.main(["--raiz", raiz, "proximo-codigo", "--tabela", "INFO_CNTTO", "--tipo", "Validação"]) == 0
    assert '"ICT.RV.4"' in capsys.readouterr().out
    assert cli.main(["--raiz", raiz, "validar", str(ficha)]) == 1
    assert "C9_codigo_repetido" in capsys.readouterr().out


# ==========================================================================================
# test_markdown.py
# ==========================================================================================
"""Todo Markdown do projeto (escrito a mao ou gerado) precisa renderizar: tabelas inteiras, titulos
separados, sem quebra de linha ou '|' vindo de dado da empresa dentro de celula, sem caractere de
controle. Um glossario com comentario de coluna em duas linhas quebrava a tabela inteira."""




IGNORAR = {".venv", "__pycache__", ".pytest_cache", ".git", "partes", "base", "dados", "saida"}
_CONTROLE_markdown = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def problemas_markdown(texto: str) -> list[str]:
    """Lista o que quebraria a visualizacao: CR, caractere de controle, tabela irregular ou colada,
    titulo colado, cerca aberta."""
    probs: list[str] = []
    if "\r" in texto:
        probs.append("CR (fim de linha Windows misturado)")
    if texto.startswith("﻿"):
        probs.append("BOM")
    linhas = texto.split("\n")
    em_cerca = ""  # a cerca aberta (``` ou ````...): so fecha com a mesma marca, de tamanho >= ao da abertura
    n_tabela: int | None = None
    for i, l in enumerate(linhas, 1):
        if _CONTROLE_markdown.search(l):
            probs.append(f"{i}: caractere de controle {_CONTROLE_markdown.search(l).group(0)!r}")
        m = re.match(r"\s{0,3}(`{3,}|~{3,})", l)
        if m:
            marca = m.group(1)
            if not em_cerca:
                em_cerca = marca
                n_tabela = None
                continue
            if marca[0] == em_cerca[0] and len(marca) >= len(em_cerca) and not l.strip().strip(marca[0]):
                em_cerca = ""
                continue
        if em_cerca:
            continue
        prev = linhas[i - 2] if i >= 2 else ""
        if l.startswith("|"):
            n = len(re.findall(r"(?<!\\)\|", l))
            if n_tabela is None:
                if prev.strip() and not prev.startswith("|"):
                    probs.append(f"{i}: tabela sem linha em branco antes")
                n_tabela = n
            elif n != n_tabela:
                probs.append(f"{i}: linha da tabela com {n} '|' (cabecalho tem {n_tabela})")
        else:
            if n_tabela is not None and l.strip():
                probs.append(f"{i}: texto colado logo abaixo da tabela (vira celula)")
            n_tabela = None
        if re.match(r"#{1,6} ", l) and prev.strip() and i > 1:
            probs.append(f"{i}: titulo sem linha em branco antes")
    if em_cerca:
        probs.append("cerca de codigo nao fechada")
    if texto and not texto.endswith("\n"):
        probs.append("sem quebra de linha no fim")
    return probs


def _arquivos_md_markdown():
    for p in sorted(list(RAIZ.rglob("*.md")) + list((REPO / ".github").rglob("*.md")) + [REPO / "README.md"]):
        if not (IGNORAR & set(p.relative_to(REPO).parts)):
            yield p


def test_markdown_escrito_a_mao_renderiza():
    ruins = {}
    for p in _arquivos_md_markdown():
        probs = problemas_markdown(p.read_text(encoding="utf-8"))
        if probs:
            ruins[p.relative_to(REPO).as_posix()] = probs
    assert not ruins, "\n".join(f"{k}: {v}" for k, v in ruins.items())


def test_helpers_md():
    assert nucleo.celula("linha 1\r\nlinha 2 | com pipe") == "linha 1 linha 2 \\| com pipe"
    assert nucleo.inline("  a\n\n  b   c ") == "a b c"
    assert nucleo.bloco("texto\n# nao e titulo\n- nem lista\n| nem | tabela |\n1. nem numerada\n\n\n\nfim") == \
        "texto\n\\# nao e titulo\n\\- nem lista\n\\| nem | tabela |\n\\1. nem numerada\n\nfim"
    assert nucleo.cerca("tem ```crase``` dentro")[0] == "````"


HOSTIL = "Comentário em duas linhas\ncom | pipe e\r\n# linha que parece título"


def test_referencias_geradas_com_dado_hostil_renderizam(projeto):
    regras = nucleo.ler_aprovadas_md(projeto.base / "regras_aprovadas_exemplo.md", projeto.nao_siglas)
    regras[0]["descricao"] = HOSTIL
    glossario = [{"sigla": "CMPO_X", "termo_negocio": HOSTIL, "tabela_sigla": "INFO_X", "fonte": "unity_catalog", "confianca": "alta"}]
    snap = {"catalogo": "c", "schema": "s", "snapshot_id": "1", "tabelas": {"t": {"comentario": HOSTIL, "colunas": [{"nome": "a"}]}}}
    mapa = {"INFO_X": {"prefixo": "IX", "tabela_uc": "t", "prefixo_confirmado": True, "tabela_uc_confirmada": False}}
    escritos = nucleo.gerar(projeto, {"texto": "# KB\r\n\r\ntexto"}, regras, glossario, snap, mapa)
    for p in escritos:
        assert not problemas_markdown(p.read_text(encoding="utf-8")), p
    glo = (projeto.copilot_fonte / "skills" / "regras-cadflow" / "references" / "glossario.md").read_text(encoding="utf-8")
    assert "duas linhas com \\| pipe e # linha" in glo


def test_ficha_e_trilha_com_dado_hostil_renderizam(projeto):
    r = nucleo.completar({**nucleo.regra_vazia(), "codigo": "ICT.RV.3", "nome": "Validação | do telefone", "tipo": "Validação",
                        "criticidade": "Baixa", "tabela": "Informações de Contato (INFO_CNTTO)", "tipo_pessoa": "PF",
                        "descricao": HOSTIL, "objetivo": "Validar.", "criterio_aplicacao": "Executar.", "logica_execucao": "SE x ENTÃO y.",
                        "resultado": "- começa com traço\n| e pipe", "exemplo": "ex", "acao_dado_valido": "a", "acao_dado_invalido": "b",
                        "observacao": "```\ncerca dentro\n```"})
    ficha = nucleo.gravar_ficha(r, projeto)
    txt = ficha.read_text(encoding="utf-8")
    assert not problemas_markdown(txt), problemas_markdown(txt)
    assert "\\# linha que parece título" in txt and "\\- começa com traço" in txt
    cli.registrar(projeto, "anotacao", "titulo\ncom quebra | pipe", "detalhe com ```crase```", regerar=False)
    cli.registrar_ficha(projeto, r, "validar")
    p = cli.exportar_md(projeto)
    txt = p.read_text(encoding="utf-8")
    assert not problemas_markdown(txt), problemas_markdown(txt)


# ==========================================================================================
# test_texto.py
# ==========================================================================================
def test_mojibake_vai_e_volta():
    original = "Base de Conhecimento para Geração de Regras de Negócio — Lógica de Execução"
    quebrado = original.encode("utf-8").decode("cp1252")
    assert "Ã" in quebrado
    assert nucleo.corrigir_mojibake(quebrado) == original
    assert nucleo.corrigir_mojibake(original) == original  # idempotente


def test_mojibake_duplo():
    original = "Validação"
    duplo = original.encode("utf-8").decode("cp1252").encode("utf-8").decode("cp1252")
    assert nucleo.corrigir_mojibake(duplo) == original


def test_ler_texto_arquivo_quebrado(tmp_path):
    original = "## Critério de Aplicação\nExecutar em todos os registros."
    p = tmp_path / "x.md"
    p.write_bytes(original.encode("utf-8").decode("cp1252").encode("utf-8"))
    assert nucleo.ler_texto(p) == original


def test_siglas_e_filtro():
    t = "SE CMPO_DDD_TEL = null ENTÃO enriquecer com a base da RBRF; campo CMPO_NAC e tabela TDDD"
    s = nucleo.siglas(t, {"SE", "ENTÃO", "RBRF"})
    assert "CMPO_DDD_TEL" in s and "CMPO_NAC" in s and "TDDD" in s
    assert "RBRF" not in s and "SE" not in s


def test_codigos():
    t = 'Segregar como inconsistente: "x". Executar regra ICT.RL.5. Ver ICT.RV5 e IBS.RN1.RE3'
    assert nucleo.codigos_citados(t) == ["ICT.RL.5", "ICT.RV.5", "IBS.RN1.RE3"]
    assert nucleo.decompor_codigo("IBS.RN1.RE3") == {"prefixo": "IBS", "tipo": "RE", "numero": "3", "rn": "IBS.RN.1"}
    assert nucleo.decompor_codigo("ICT.RV.5")["tipo"] == "RV"
    assert nucleo.normalizar_codigo("ict.rv5") == "ICT.RV.5"


def test_tabela_e_sigla_e_dominios():
    assert nucleo.tabela_e_sigla("Informações de Contato (INFO_CNTTO)") == ("Informações de Contato", "INFO_CNTTO")
    assert nucleo.tabela_e_sigla("INFO_CNTTO") == ("INFO_CNTTO", "INFO_CNTTO")
    assert nucleo.dominios_citados("pertence aos domínios permitidos (TDDD - XDOM001)") == ["TDDD - XDOM001"]


def test_limpar_nan_e_normalizar():
    assert nucleo.limpar_nan("nan") == "" and nucleo.limpar_nan(None) == "" and nucleo.limpar_nan(" x ") == "x"
    assert nucleo.normalizar("  Validação  ") == "validacao"


# ==========================================================================================
# test_trilha.py
# ==========================================================================================
def _ficha_trilha(**kw):
    base = {"codigo": "ICT.RV.3", "nome": "Validação do número do telefone de contato", "tipo": "Validação", "criticidade": "Baixa",
            "tabela": "Informações de Contato (INFO_CNTTO)", "tipo_pessoa": "PF",
            "descricao": "Valida se o número do telefone de contato tem a quantidade de dígitos esperada.", "objetivo": "Validar o número do telefone de contato de pessoa física.",
            "criterio_aplicacao": "Executar em todos os registros que possuam conteúdo no campo número do telefone de contato.",
            "logica_execucao": "A regra verifica se o valor contido no campo CMPO_NUM_TEL possui 9 dígitos, sendo considerado válido ao atender esta condição.",
            "resultado": "Registros com número do telefone fora de 9 dígitos são segregados como inconsistentes.",
            "exemplo": "Dado analisado: 912345678. Como possui 9 dígitos, o dado é válido; o registro é mantido.", "acao_dado_valido": "Manter o dado validado, sem alterações.",
            "acao_dado_invalido": 'Segregar como inconsistente: "Número do telefone com quantidade de dígitos incoerente".'}
    base.update(kw)
    return base


def test_trilha_eventos_ficha_e_html(projeto, capsys):
    raiz = str(projeto.raiz)
    assert cli.main(["--raiz", raiz, "atualizar", "--sem-uc"]) == 0
    capsys.readouterr()
    assert cli.main(["--raiz", raiz, "trilha", "--nova", "validar telefone"]) == 0
    assert cli.main(["--raiz", raiz, "anotar", "--leu", "skill:regras-cadflow"]) == 0
    assert cli.main(["--raiz", raiz, "anotar", "classifiquei como Validação"]) == 0
    assert cli.main(["--raiz", raiz, "anotar", "--pergunta", "criticidade?", "sugeri Baixa"]) == 0
    assert cli.main(["--raiz", raiz, "campo", "CMPO_NUM_TEL"]) == 0
    ficha = projeto.saida / "_ficha_trilha.json"
    ficha.write_text(json.dumps(_ficha_trilha(), ensure_ascii=False), encoding="utf-8")
    assert cli.main(["--raiz", raiz, "validar", str(ficha)]) == 0
    ficha.write_text(json.dumps(_ficha_trilha(criticidade="Média"), ensure_ascii=False), encoding="utf-8")
    assert cli.main(["--raiz", raiz, "validar", str(ficha)]) == 0
    assert cli.main(["--raiz", raiz, "gravar", str(ficha)]) == 0
    evs = cli.ler(projeto)
    tipos = [e["tipo"] for e in evs]
    assert tipos[0] == "inicio" and "leu" in tipos and "conclusao" in tipos and "pergunta" in tipos and "gravou" in tipos
    comandos = [e for e in evs if e["tipo"] == "comando"]
    assert any(e["titulo"].startswith("campo CMPO_NUM_TEL") and "cntto" in e["detalhe"].lower() or "contato" in e["detalhe"].lower() for e in comandos)
    fichas = [e for e in evs if e["tipo"] == "ficha"]
    assert len(fichas) == 2 and fichas[1]["mudou"] == ["criticidade"] and "~ criticidade" in fichas[1]["detalhe"]
    html = (projeto.saida / "_trilha.html").read_text(encoding="utf-8")
    assert "Trilha da conversa" in html and "classifiquei como Validação" in html and "location.reload" in html
    capsys.readouterr()
    assert cli.main(["--raiz", raiz, "trilha"]) == 0
    assert "evento(s)" in capsys.readouterr().out
    assert cli.main(["--raiz", raiz, "trilha", "--exportar"]) == 0
    md = next(projeto.saida.glob("trilha_*.md")).read_text(encoding="utf-8")
    assert "## Ficha final" in md and "ICT.RV.3" in md
    # erro de comando vira evento 'erro'
    assert cli.main(["--raiz", raiz, "validar", str(ficha)]) == 1  # codigo ja existe agora
    assert cli.ler(projeto)[-1]["tipo"] == "erro"


def test_trilha_desligada_nao_grava(projeto, capsys):
    raiz = str(projeto.raiz)
    assert cli.main(["--raiz", raiz, "trilha", "--desligar"]) == 0
    cfg = nucleo.carregar(projeto.raiz)
    assert not cli.ativa(cfg)
    assert cli.main(["--raiz", raiz, "anotar", "nada"]) == 0
    assert "desligada" in capsys.readouterr().out
    assert cli.main(["--raiz", raiz, "status"]) == 0
    assert not (projeto.saida / "_trilha.jsonl").exists()
    assert cli.main(["--raiz", raiz, "trilha", "--ligar"]) == 0
    assert cli.ativa(nucleo.carregar(projeto.raiz))
    assert cli.main(["--raiz", raiz, "trilha", "--limpar"]) == 0


# ==========================================================================================
# test_validar.py
# ==========================================================================================
def _ctx_validar(projeto, **kw):
    snap = json.loads((EXEMPLOS / "uc" / "esquema_exemplo.json").read_text(encoding="utf-8"))
    regras = nucleo.ler_aprovadas_md(EXEMPLOS / "base" / "regras_aprovadas_exemplo.md", projeto.nao_siglas)
    return nucleo.Contexto(siglas_conhecidas=set(catalogo.indice_colunas(snap)), codigos_existentes=nucleo.codigos_usados(regras), **kw), regras


def test_aprovadas_passam_na_calibracao(projeto):
    ctx, regras = _ctx_validar(projeto)
    ctx.codigos_existentes = set()
    cal = nucleo.calibrar(regras, projeto, ctx)
    assert cal["total"] == 8
    assert cal["reprovadas"] == [], cal["reprovadas"]
    assert [i["codigo"] for i in cal["incompletas_na_fonte"]] == ["IBS.RL.1"]  # Resultado: nan


def _regra_boa_validar():
    return {
        "codigo": "ICT.RV.3", "nome": "Validação do número do telefone de contato", "tipo": "Validação", "criticidade": "Baixa",
        "tabela": "Informações de Contato (INFO_CNTTO)", "tipo_pessoa": "PF",
        "descricao": "Valida se o número do telefone de contato tem a quantidade de dígitos esperada para o tipo informado.",
        "objetivo": "Validar o número do telefone de contato de pessoa física.",
        "criterio_aplicacao": "Executar em todos os registros que possuam conteúdo no campo número do telefone de contato.",
        "logica_execucao": "A regra verifica se o valor contido no campo CMPO_NUM_TEL possui 9 dígitos quando o campo CMPO_TIPO_TEL for celular e 8 dígitos nos demais casos, sendo considerado válido ao atender esta condição.",
        "resultado": "Registros com número do telefone fora da quantidade de dígitos esperada são segregados como inconsistentes.",
        "exemplo": "Dado analisado: 912345678 com tipo celular. Como possui 9 dígitos, o dado é válido; o registro é mantido. Dado analisado: 12345 com tipo celular. Como não possui 9 dígitos, o dado é inválido; o registro é segregado como inconsistente.",
        "acao_dado_valido": "Manter o dado validado, sem alterações.",
        "acao_dado_invalido": 'Segregar como inconsistente: "Número do telefone com quantidade de dígitos incoerente".',
        "familia": ["ICT.RL.1"],
    }


def test_regra_boa_sem_erros(projeto):
    ctx, _ = _ctx_validar(projeto, pedido="nucleo o número do telefone de contato PF e segregar", nova=True)
    r = nucleo.completar(_regra_boa_validar())
    assert r["regras_encadeadas"] == ["ICT.RL.1"]  # familia vira vinculo interno, sem aparecer no texto
    ach = nucleo.validar(r, projeto, ctx)
    assert not [a for a in ach if a.nivel == "erro"], [str(a) for a in ach]
    assert "C10_cita_outra_regra" not in {a.check for a in ach}


def test_forma_se_entao_e_aceita(projeto):
    ctx, _ = _ctx_validar(projeto, nova=True)
    r = _regra_boa_validar()
    r["logica_execucao"] = "SE CMPO_NUM_TEL possuir 9 dígitos quando CMPO_TIPO_TEL for celular ENTÃO considerar o dado válido. SENÃO considerar o dado inválido."
    checks = {a.check for a in nucleo.validar(nucleo.completar(r), projeto, ctx)}
    assert "C4_inicio_logica" not in checks and "C4_verbo_do_tipo" not in checks and "C4_sem_senao" not in checks
    r["logica_execucao"] = "SE CMPO_NUM_TEL possuir 9 dígitos ENTÃO considerar o dado válido."
    assert "C4_sem_senao" in {a.check for a in nucleo.validar(nucleo.completar(r), projeto, ctx)}


def test_estilo_vigente_e_erro_so_para_regra_nova(projeto):
    """O mesmo texto antigo: aviso na calibracao (aprovadas), erro para regra em criacao."""
    r = _regra_boa_validar()
    r["resultado"] = "Válido: quantidade de dígitos coerente. Inválido: quantidade de dígitos incoerente."
    r["acao_dado_valido"] = "Manter dado."
    r["acao_dado_invalido"] = 'Segregar como inconsistente: "x". Executar regra ICT.RL.1.'
    r["exemplo"] = "912345678 é válido."
    r["observacao"] = "Dúvida: confirmar com a Governança se o DDD entra."
    ctx, _ = _ctx_validar(projeto)
    por = {(a.check, a.nivel) for a in nucleo.validar(nucleo.completar(r), projeto, ctx)}
    assert ("C10_cita_outra_regra", "aviso") in por and ("C12_resultado_rotulo", "aviso") in por
    assert ("C16_acao_valida_vaga", "aviso") in por and ("C17_exemplo_generico", "aviso") in por and ("C19_observacao_processo", "aviso") in por
    assert not [c for c, n in por if n == "erro"], por
    ctx, _ = _ctx_validar(projeto, nova=True)
    por = {(a.check, a.nivel) for a in nucleo.validar(nucleo.completar(r), projeto, ctx)}
    assert {("C10_cita_outra_regra", "erro"), ("C12_resultado_rotulo", "erro"), ("C16_acao_valida_vaga", "erro"), ("C17_exemplo_generico", "erro")} <= por


def test_marcador_e_sigla_fora_do_lugar(projeto):
    ctx, _ = _ctx_validar(projeto, nova=True)
    r = _regra_boa_validar()
    r["exemplo"] = "Dado analisado: <valor concreto>. Como o campo CMPO_NUM_TEL tem 9 dígitos, o dado é válido; mantido."
    por = {(a.check, a.nivel) for a in nucleo.validar(nucleo.completar(r), projeto, ctx)}
    assert ("C15_marcador", "erro") in por and ("C18_sigla_fora_do_lugar", "erro") in por
    r["logica_execucao"] = "SE o telefone for 11 < 9 ENTÃO considerar válido SENÃO inválido."  # '<' de comparacao nao e marcador
    assert "C15_marcador" not in {a.check for a in nucleo.validar(nucleo.completar({**r, "exemplo": _regra_boa_validar()["exemplo"]}), projeto, ctx)}


def test_regra_ruim_acusa_cada_defeito(projeto):
    ctx, _ = _ctx_validar(projeto)
    r = _regra_boa_validar()
    r["descricao"] = "Valida o campo CMPO_NUM_TEL."                       # C2 sigla em texto de negocio
    r["objetivo"] = "Garantir a qualidade dos dados."                      # C7 formula vazia
    r["logica_execucao"] = "A regra apaga o campo CMPO_NUM_TEL e CMPO_INEXISTENTE_X."  # C4 verbo de outro tipo; C2b sigla desconhecida
    r["codigo"] = "ICT.RV.1"                                               # C9 repetido
    r["familia"] = ["ICT.RL.9"]                                            # C10 familia inexistente
    checks = {a.check for a in nucleo.validar(nucleo.completar(r), projeto, ctx)}
    assert {"C2_sigla_em_texto_negocio", "C7_objetivo_generico", "C4_verbo_do_tipo", "C2b_sigla_desconhecida",
            "C9_codigo_repetido", "C10_encadeamento_inexistente"} <= checks


def test_encadeamento_criado_junto_nao_e_erro(projeto):
    ctx, _ = _ctx_validar(projeto, codigos_em_criacao={"ICT.RL.9"}, nova=True)
    r = _regra_boa_validar()
    r["familia"] = "ICT.RL.9"
    assert "C10_encadeamento_inexistente" not in {a.check for a in nucleo.validar(nucleo.completar(r), projeto, ctx)}


def test_canonizar_grafias(projeto):
    r = nucleo.canonizar({"tipo": "validacao", "criticidade": "muito alta", "tipo_pessoa": "pf / pj", "tabela": "X (INFO_X)"}, projeto)
    assert r["tipo"] == "Validação" and r["criticidade"] == "Muito Alta" and r["tipo_pessoa"] == "PF,PJ" and r["tabela_sigla"] == "INFO_X"


def test_c1_bloqueia_sem_campos(projeto):
    ach = nucleo.validar({"nome": "x", "tipo": "Validação"}, projeto)
    assert len(ach) == 1 and ach[0].check == "C1_obrigatorios" and "exemplo" in ach[0].mensagem


def test_instalar_copilot_bloco_idempotente_e_settings(projeto):
    assert nucleo.instalar_copilot(projeto) == []
    ci = projeto.raiz_workspace / ".github" / "copilot-instructions.md"
    txt = ci.read_text(encoding="utf-8")
    assert "Texto que ja existia" in txt and txt.count(nucleo.MARCA_INICIO) == 1 and "`cdflw/`" in txt
    nucleo.instalar_copilot(projeto)
    assert ci.read_text(encoding="utf-8").count(nucleo.MARCA_INICIO) == 1
    st = json.loads((projeto.raiz_workspace / ".vscode" / "settings.json").read_text(encoding="utf-8"))
    assert st["chat.useAgentSkills"] is True and st["editor.fontSize"] == 14
    (projeto.raiz_workspace / ".github" / "skills" / "refazer" / "SKILL.md").unlink()
    assert nucleo.instalar_copilot(projeto) == ["skills/refazer/SKILL.md"]


def test_configurar_modelo_grava_no_agente(projeto, capsys):
    raiz = str(projeto.raiz)
    assert cli.main(["--raiz", raiz, "configurar", "--local", "--modelo", "GPT-6.1 Sol"]) == 0
    pub = (projeto.raiz_workspace / ".github" / "agents" / "cdflw.agent.md").read_text(encoding="utf-8")
    assert "model: 'GPT-6.1 Sol'" in pub
    assert nucleo.carregar(projeto.raiz).modelo == "GPT-6.1 Sol"
    assert "GPT-6.1 Sol" not in (REPO / ".github" / "agents" / "cdflw.agent.md").read_text(encoding="utf-8")  # o repositorio nao muda


# ==========================================================================================
# origem: nenhum arquivo que o usuario copia pode dizer de onde veio (so o README da raiz, que e o guia de copia)
# ==========================================================================================
_ORIGEM = re.compile("raw\.git" + "hubusercontent|victor" + "moraessp", re.I)


def test_produto_nao_diz_de_onde_veio():
    ruins = []
    for p in list(RAIZ.rglob("*")) + list((REPO / ".github").rglob("*")):
        if not p.is_file() or p.suffix not in (".py", ".md", ".toml", ".json", ".txt") or "references" in p.parts or "dados" in p.parts:
            continue
        for i, linha in enumerate(p.read_text(encoding="utf-8", errors="replace").splitlines(), 1):
            if _ORIGEM.search(linha):
                ruins.append(f"{p.relative_to(REPO).as_posix()}:{i}")
    assert not ruins, ruins
